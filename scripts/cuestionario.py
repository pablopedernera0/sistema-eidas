#!/usr/bin/env python3
"""
Cuestionarios de opción múltiple a partir de un único archivo markdown.

El banco de preguntas (con la clave) vive en sistema-eidas-datos, que es privado:
este repo es público y nunca tiene que guardar la clave.

Formato del archivo de preguntas:

    ---
    id: simulacro-2p-bicirio
    titulo: ...
    semilla: 2026
    ---

    ### 1. (2 pts) Enunciado
    - [ ] opción incorrecta
    - [x] opción correcta
    - [ ] ...
    > Por qué: explicación que ven al corregir.

Comandos:

    generar  <preguntas.md> --sitio <carpeta> --privado <carpeta>
        Escribe <sitio>/preguntas.js (sin clave, se publica ya) y
        <privado>/clave.js (se publica recién cuando termina el cuestionario).
        El orden de las opciones se mezcla con la semilla, siempre igual.

    corregir <preguntas.md> <archivos de respuestas...> [--csv salida.csv]
        Lee los .txt que generan los alumnos al finalizar, calcula el puntaje y
        muestra qué preguntas fallaron más.

    kahoot   <preguntas.md> [-o salida.xlsx]
        Arma la planilla para "Importar hoja de cálculo" de Kahoot (mismo formato
        que la plantilla oficial). Valida los límites de Kahoot: 120 caracteres la
        pregunta, 75 cada respuesta, hasta 4 respuestas. Tiempo por pregunta con
        "### N. (30 s) ..." o `tiempo:` en el frontmatter (20 s si no hay nada).
        Admite más de una [x].
"""
import argparse
import csv
import json
import random
import re
import sys
from pathlib import Path


def leer_frontmatter(texto):
    """Frontmatter simple: `clave: valor` y bloques `clave: |` indentados."""
    meta, cuerpo = {}, texto
    if texto.startswith('---\n'):
        fin = texto.index('\n---\n', 4)
        bloque, cuerpo = texto[4:fin], texto[fin + 5:]
        clave_bloque = None
        for linea in bloque.split('\n'):
            if clave_bloque and (linea.startswith('  ') or not linea.strip()):
                meta[clave_bloque] += linea[2:] + '\n'
                continue
            clave_bloque = None
            if ':' in linea:
                k, v = linea.split(':', 1)
                k, v = k.strip(), v.strip()
                if v == '|':
                    meta[k], clave_bloque = '', k
                else:
                    meta[k] = v
        meta = {k: v.strip() if isinstance(v, str) else v for k, v in meta.items()}
    return meta, cuerpo


def leer_preguntas(ruta):
    meta, cuerpo = leer_frontmatter(Path(ruta).read_text(encoding='utf-8'))
    preguntas = []
    for bloque in re.split(r'^### ', cuerpo, flags=re.M)[1:]:
        lineas = bloque.strip().split('\n')
        m = re.match(r'(\d+)\.\s*(?:\((\d+)\s*pts?\)\s*)?(?:\((\d+)\s*s\)\s*)?(.*)', lineas[0])
        if not m:
            sys.exit(f'Encabezado de pregunta mal formado: {lineas[0]!r}')
        num, pts, enunciado = int(m[1]), int(m[2] or 1), m[4].strip()
        tiempo = int(m[3]) if m[3] else None
        opciones, correctas, porque = [], [], ''
        for linea in lineas[1:]:
            op = re.match(r'- \[( |x)\] (.*)', linea)
            if op:
                if op[1] == 'x':
                    correctas.append(len(opciones))
                opciones.append(op[2].strip())
            elif linea.startswith('> Por qué:'):
                porque = linea[len('> Por qué:'):].strip()
        if not correctas or len(opciones) < 2:
            sys.exit(f'Pregunta {num}: tiene que tener opciones y al menos una [x].')
        preguntas.append(dict(num=num, puntos=pts, enunciado=enunciado, tiempo=tiempo,
                              opciones=opciones, correcta=correctas[0], correctas=correctas,
                              porque=porque))
    if not preguntas:
        sys.exit('No encontré preguntas (encabezados "### N. ...").')
    return meta, preguntas


def mezclar(meta, preguntas):
    """Mezcla las opciones de cada pregunta de forma reproducible."""
    semilla = int(meta.get('semilla', 0))
    letras = 'abcdefgh'
    publicas, clave = [], {}
    for p in preguntas:
        orden = list(range(len(p['opciones'])))
        random.Random(semilla * 1000 + p['num']).shuffle(orden)
        opciones = [{'letra': letras[i], 'texto': p['opciones'][j]} for i, j in enumerate(orden)]
        letras_ok = [letras[orden.index(c)] for c in p['correctas']]
        publicas.append(dict(num=p['num'], puntos=p['puntos'],
                             enunciado=p['enunciado'], opciones=opciones))
        clave[str(p['num'])] = dict(correcta=letras_ok[0], porque=p['porque'])
        if len(letras_ok) > 1:
            clave[str(p['num'])]['correctas'] = letras_ok
    return publicas, clave


def cuestionario_id(meta, ruta):
    return meta.get('id') or Path(ruta).stem


def generar(args):
    meta, preguntas = leer_preguntas(args.preguntas)
    publicas, clave = mezclar(meta, preguntas)
    qid = cuestionario_id(meta, args.preguntas)
    total = sum(p['puntos'] for p in preguntas)
    datos = dict(id=qid, titulo=meta.get('titulo', qid), total=total, preguntas=publicas)
    sitio, privado = Path(args.sitio), Path(args.privado)
    sitio.mkdir(parents=True, exist_ok=True)
    privado.mkdir(parents=True, exist_ok=True)
    (sitio / 'preguntas.js').write_text(
        '// Generado por sistema-eidas/scripts/cuestionario.py. No editar a mano.\n'
        'window.CUESTIONARIO = ' + json.dumps(datos, ensure_ascii=False, indent=1) + ';\n',
        encoding='utf-8')
    (privado / 'clave.js').write_text(
        '// Generado por sistema-eidas/scripts/cuestionario.py. Publicar solo al terminar.\n'
        'window.CLAVE = ' + json.dumps(dict(id=qid, respuestas=clave), ensure_ascii=False, indent=1)
        + ';\n', encoding='utf-8')
    reparto = {}
    for v in clave.values():
        reparto[v['correcta']] = reparto.get(v['correcta'], 0) + 1
    print(f'{len(preguntas)} preguntas, {total} puntos. Correctas por letra: {dict(sorted(reparto.items()))}')
    print(f'  público : {sitio / "preguntas.js"}')
    print(f'  privado : {privado / "clave.js"}  (subir al sitio recién al terminar)')


def leer_respuesta(ruta):
    texto = Path(ruta).read_text(encoding='utf-8', errors='replace')
    m = re.search(r'^DATOS: (\{.*\})\s*$', texto, flags=re.M)
    if not m:
        raise ValueError('no tiene la línea DATOS')
    return json.loads(m[1])


def corregir(args):
    meta, preguntas = leer_preguntas(args.preguntas)
    _, clave = mezclar(meta, preguntas)
    qid = cuestionario_id(meta, args.preguntas)
    total = sum(p['puntos'] for p in preguntas)
    filas, aciertos, vistos = [], {p['num']: 0 for p in preguntas}, set()
    for ruta in args.respuestas:
        try:
            d = leer_respuesta(ruta)
        except (ValueError, json.JSONDecodeError) as e:
            print(f'  salteado {ruta}: {e}', file=sys.stderr)
            continue
        if d.get('id') != qid:
            print(f'  salteado {ruta}: es de otro cuestionario ({d.get("id")})', file=sys.stderr)
            continue
        clave_alumno = (d.get('nombre', '').strip().lower(), d.get('fin'))
        if clave_alumno in vistos:
            continue  # el mismo archivo mandado dos veces
        vistos.add(clave_alumno)
        r = d.get('r', {})
        puntos, marcas = 0, []
        for p in preguntas:
            dada = r.get(str(p['num']), '')
            c = clave[str(p['num'])]
            ok = dada in c.get('correctas', [c['correcta']])
            puntos += p['puntos'] if ok else 0
            aciertos[p['num']] += ok
            marcas.append('ok' if ok else (dada or '-'))
        filas.append([d.get('nombre', '?'), d.get('comision', ''), d.get('fin', ''), puntos] + marcas)

    if not filas:
        sys.exit('No hubo archivos válidos para corregir.')
    filas.sort(key=lambda f: (f[1], f[0].lower()))
    print(f'\n{len(filas)} respuestas — {meta.get("titulo", qid)}\n')
    for f in filas:
        print(f'  {f[3]:>3}/{total}  {f[1]:<5} {f[0]}')
    promedio = sum(f[3] for f in filas) / len(filas)
    print(f'\n  Promedio: {promedio:.1f}/{total}\n\nPreguntas con más errores:')
    for num, n in sorted(aciertos.items(), key=lambda kv: kv[1])[:8]:
        print(f'  P{num:<3} {n}/{len(filas)} bien ({100 * n // len(filas)}%)')
    if args.csv:
        with open(args.csv, 'w', newline='', encoding='utf-8') as fh:
            w = csv.writer(fh)
            w.writerow(['nombre', 'comision', 'finalizado', f'puntaje/{total}']
                       + [f'P{p["num"]}' for p in preguntas])
            w.writerows(filas)
        print(f'\nPlanilla: {args.csv}')


TIEMPOS_KAHOOT = (5, 10, 20, 30, 60, 90, 120, 240)


def kahoot(args):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    meta, preguntas = leer_preguntas(args.preguntas)
    semilla = int(meta.get('semilla', 0))
    tiempo_base = int(meta.get('tiempo', 20))
    errores, filas = [], []
    for p in preguntas:
        tiempo = p['tiempo'] or tiempo_base
        if len(p['enunciado']) > 120:
            errores.append(f'P{p["num"]}: la pregunta tiene {len(p["enunciado"])} caracteres (máx. 120)')
        if len(p['opciones']) > 4:
            errores.append(f'P{p["num"]}: Kahoot admite hasta 4 respuestas')
        for o in p['opciones']:
            if len(o) > 75:
                errores.append(f'P{p["num"]}: respuesta de {len(o)} caracteres (máx. 75): {o[:40]}...')
        if tiempo not in TIEMPOS_KAHOOT:
            errores.append(f'P{p["num"]}: tiempo {tiempo} s no válido, usar {TIEMPOS_KAHOOT}')
        orden = list(range(len(p['opciones'])))
        random.Random(semilla * 1000 + p['num']).shuffle(orden)
        opciones = [p['opciones'][j] for j in orden]
        correctas = ','.join(str(orden.index(c) + 1) for c in p['correctas'])
        filas.append((p['enunciado'], opciones, tiempo, correctas))
    if errores:
        sys.exit('No se generó la planilla:\n  ' + '\n  '.join(errores))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws['B2'] = 'Quiz template'
    ws['B2'].font = Font(bold=True, size=14)
    ws['B3'] = meta.get('titulo', '')
    encabezados = ['Question - max 120 characters', 'Answer 1 - max 75 characters',
                   'Answer 2 - max 75 characters', 'Answer 3 - max 75 characters',
                   'Answer 4 - max 75 characters',
                   'Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs',
                   'Correct answer(s) - choose at least one']
    relleno = PatternFill('solid', fgColor='46178F')
    for i, h in enumerate(encabezados):
        celda = ws.cell(row=8, column=2 + i, value=h)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = relleno
        celda.alignment = Alignment(wrap_text=True, vertical='center')
    for r, (q, ops, tiempo, correctas) in enumerate(filas, start=9):
        ws.cell(row=r, column=1, value=r - 8)
        ws.cell(row=r, column=2, value=q)
        for j, o in enumerate(ops):
            ws.cell(row=r, column=3 + j, value=o)
        ws.cell(row=r, column=7, value=tiempo)
        ws.cell(row=r, column=8, value=correctas)
        for col in range(2, 9):
            ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical='top')
    for col, ancho in zip('ABCDEFGH', [5, 50, 32, 32, 32, 32, 14, 14]):
        ws.column_dimensions[col].width = ancho
    salida = args.o or str(Path(args.preguntas).with_suffix('.xlsx'))
    wb.save(salida)
    reparto = {}
    for *_, c in filas:
        reparto[c] = reparto.get(c, 0) + 1
    duracion = sum(f[2] for f in filas)
    print(f'{len(filas)} preguntas, ~{duracion // 60 + len(filas) // 2} min de juego. '
          f'Correcta en la posición: {dict(sorted(reparto.items()))}')
    print(f'  planilla: {salida}')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    g = sub.add_parser('generar')
    g.add_argument('preguntas')
    g.add_argument('--sitio', required=True, help='carpeta pública (github.io)')
    g.add_argument('--privado', required=True, help='carpeta privada para clave.js')
    c = sub.add_parser('corregir')
    c.add_argument('preguntas')
    c.add_argument('respuestas', nargs='+')
    c.add_argument('--csv')
    k = sub.add_parser('kahoot')
    k.add_argument('preguntas')
    k.add_argument('-o', help='planilla de salida (por defecto, al lado del .md)')
    args = ap.parse_args()
    {'generar': generar, 'corregir': corregir, 'kahoot': kahoot}[args.cmd](args)


if __name__ == '__main__':
    main()
